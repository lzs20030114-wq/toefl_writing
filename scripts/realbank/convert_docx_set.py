#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""把「docx 套壳的截图套题」转换成旧管线（ingest_set.py）认识的布局。

背景：旧源套题是 `写作.pdf 口语.pdf 听力.pdf 答案.pdf 阅读.pdf` + 音频；
5 月起新收的 7 套真新题格式相同但载体是 docx——正文四科（写作/口语/听力/阅读）
是按文档顺序贴的截图，只有「答案」docx 是纯文本。`ingest_set.py` 只认 .pdf
（非 PDF 一律 `role=None, note="非 PDF，本阶段跳过"`），所以这里把每份 docx
转成一份同名 PDF：
  - 四科正文：把 inline_shapes 按文档顺序抽出来，每张图一页，交给下游 OCR
    （本地 RapidOCR，零 token，见 scripts/ops/ocr_common.py）。
  - 答案：把段落文本用 PyMuPDF 内置 CJK 字体("china-ss") 写成带文字层的 PDF，
    这样 ingest_set.py 的 `pdf_text()` 直接走 text-layer 分支，连 OCR 都不用跑，
    比图片答案页更干净。
音频 zip 解压出 m4a（Mac 压的包文件名常被 cp437 错编码，这里按
`encode('cp437').decode('utf-8')` 修回来，__MACOSX/ 与 `._*` 垃圾一律丢弃），
连同已解压的 m4a 一起摊平进输出目录。

绝不写回桌面源目录；输出固定在 `.codex-tmp/realbank/src-converted/<规范化卷名>/`。
与桌面已有旧卷同名的一律加 `_v2` 后缀，避免任何人误把它当成覆盖旧卷。

用法:
  python scripts/realbank/convert_docx_set.py 5.10                 # 转一套
  python scripts/realbank/convert_docx_set.py --all                # 转全部 7 套
  python scripts/realbank/convert_docx_set.py 5.10 --json          # 附带落盘转换报告
  python scripts/realbank/convert_docx_set.py 5.10 --src <目录>     # 覆盖源根目录（默认桌面路径）

幂等：每次都会重新生成（覆盖同名输出文件），不会因为已存在而报错或重复追加。
"""
import argparse
import io
import json
import os
import re
import sys
import zipfile

import docx
import fitz  # PyMuPDF
from PIL import Image

DEFAULT_SRC_ROOT = r"D:\桌面\【2026改后全科真题】（持续更新中）"
# 产物相对仓库根（本文件在 <root>/scripts/realbank/ 下）：写死 D:\toefl_writing 的话
# 云端 checkout（/home/runner/work/...）会把转换结果写到一个不存在的盘符上。
REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT_ROOT = os.path.join(REPO_ROOT, ".codex-tmp", "realbank", "src-converted")

# key -> {src: 相对 DEFAULT_SRC_ROOT/--src 的子路径, out: 输出卷名}
#
# 这张表**只是 5 月那 7 套的便捷别名**，不是能转哪些卷的白名单：自动录入进来的套目录
# 事先不可能出现在这里（2.4新托福真题当初就是卡在这上面）。任意套目录走 `--dir <目录>`
# （相对 --src 根或绝对路径）+ 可选 `--out-name <卷名>`，见 convert_set_dir()。
#
# 5.10 / 5.6 在桌面根目录下已有同名旧卷（内容不同，答案键已核实不重合），
# 输出必须加 _v2 以免被误当成同一套卷覆盖。
SETS = {
    "5.10": {"src": os.path.join("5月", "5.10", "5.10 套一"), "out": "5.10新托福真题_v2"},
    "5.11": {"src": os.path.join("5月", "5.11 套一"), "out": "5.11新托福真题"},
    "5.18": {"src": os.path.join("5月", "5.18-1"), "out": "5.18新托福真题"},
    "5.20": {"src": os.path.join("5月", "5.20"), "out": "5.20新托福真题"},
    "5.23": {"src": os.path.join("5月", "5.23"), "out": "5.23新托福真题"},
    "5.29": {"src": os.path.join("5月", "5.29"), "out": "5.29新托福真题"},
    "5.6": {"src": os.path.join("5月", "5.6", "5.6 套一"), "out": "5.6新托福真题_v2"},
}

SUBJECT_KEYWORDS = ("答案", "写作", "口语", "听力", "阅读")
AUDIO_EXT = (".m4a", ".mp3", ".wav", ".mp4", ".mov")
MAX_PAGE_EDGE_PT = 1600.0  # 图片页最长边上限（pt），避免超大截图撑爆 PDF/拖慢 OCR


def subject_of(name):
    for kw in SUBJECT_KEYWORDS:
        if kw in name:
            return kw
    return None


def _blob_of_inline_shape(part, shape):
    try:
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    except AttributeError:
        return None
    rid = blip.embed
    if not rid:
        return None
    rel_part = part.related_parts.get(rid)
    return rel_part.blob if rel_part is not None else None


def extract_images_in_order(docx_path):
    """按文档正文出现顺序返回图片字节列表（inline_shapes 就是按 body 顺序遍历的）。"""
    d = docx.Document(docx_path)
    blobs = []
    for shp in d.inline_shapes:
        blob = _blob_of_inline_shape(d.part, shp)
        if blob:
            blobs.append(blob)
    return blobs


def docx_images_to_pdf(docx_path, out_path):
    """截图套壳的一科 docx → 每张图一页的 PDF。返回写入的页数。"""
    blobs = extract_images_in_order(docx_path)
    doc = fitz.open()
    for blob in blobs:
        with Image.open(io.BytesIO(blob)) as im:
            w, h = im.size
        scale = MAX_PAGE_EDGE_PT / max(w, h) if max(w, h) > MAX_PAGE_EDGE_PT else 1.0
        pw, ph = w * scale, h * scale
        page = doc.new_page(width=pw, height=ph)
        page.insert_image(fitz.Rect(0, 0, pw, ph), stream=blob)
    if doc.page_count == 0:
        doc.new_page()  # 空 docx（异常情况）也落一份空 PDF，不让下游因文件缺失崩溃
    # PyMuPDF 的 insert_image(stream=...) 默认把图片解码成未压缩位图塞进 PDF
    # （实测单张 359KB 的 RGBA 截图膨胀成 3.3MB 的裸位图，无 /Filter）；
    # 存盘时开 deflate_images 才会重新套上 Flate 压缩，体积基本回到原图量级。
    doc.save(out_path, deflate=True, deflate_images=True, garbage=4)
    doc.close()
    return len(blobs)


def docx_text_to_pdf(docx_path, out_path):
    """答案 docx（纯文本）→ 带文字层的 PDF，逐段落一个文本框，用内置 CJK 字体保留中文表头。

    ingest_common.parse_answer_pdf_with_warnings 靠整行匹配「阅读/听力/写作/口语/加试」
    表头 + 行内 `NUM_ANS` 抓编号-答案对，所以每个段落必须落在独立的文本框里，不能
    跟相邻段落粘连——那样表头行会跟下一行的答案粘在一起，直接解析不出来。
    段落内部允许折行：`Page.insert_textbox()` 会在文本框宽度内自动换行，NUM_ANS
    是逐行正则、`_is_restart` 只看题号是否回到 1，折成几行都不影响解析。
    早期用 `insert_text()` 画单行踩过坑：一份 35 题的阅读答案单行超过 250 字符，
    `insert_text()` 会在页面右边界处悄悄截断（不报错、不换行），下游解析拿到的
    只是半截答案——`insert_textbox()` 没有这个问题。
    """
    d = docx.Document(docx_path)
    lines = [p.text for p in d.paragraphs]
    doc = fitz.open()
    page_w, page_h = 595.0, 842.0  # A4
    margin, fontsize, leading = 48.0, 11.0, 16.0
    usable_w = page_w - 2 * margin
    bottom = page_h - margin

    def new_page():
        return doc.new_page(width=page_w, height=page_h)

    page = new_page()
    y = margin
    written = 0
    for line in lines:
        text = line if line.strip() else " "  # 空行也要占位，保留段落间隔给正则的锚定
        # insert_textbox() 是全有或全无：盒子不够高就整段不画，只把「溢出了多少」
        # 报成一个负数返回值，不会画一半。所以先用文字宽度粗估一个起点，再按返回值
        # 一行一行加高重试，直到 rc>=0（真正放得下）——好过猜一个「大概率够用」的
        # 倍率却在个别长行上还是不够（实测按宽度整除估的行数刚好差 1 行不够放）。
        text_w = fitz.get_text_length(text, fontname="china-ss", fontsize=fontsize)
        n_rows = max(1, int(text_w // usable_w) + 1)
        while True:
            box_h = n_rows * leading
            if y + box_h > bottom:
                page = new_page()
                y = margin
            rect = fitz.Rect(margin, y, margin + usable_w, y + box_h)
            rc = page.insert_textbox(rect, text, fontname="china-ss", fontsize=fontsize,
                                      lineheight=leading / fontsize)
            if rc >= 0:
                break
            n_rows += 1  # 没放下，加一行高度重试
        y += box_h
        written += 1
    if written == 0:
        new_page()
    doc.save(out_path, deflate=True, garbage=4)
    doc.close()
    return written


# ── 音频 ─────────────────────────────────────────────────────────────────────
def _fix_mojibake_name(name):
    """Mac 压的 zip 常把 UTF-8 文件名按 cp437 写进中央目录（flag 未置 UTF-8 位）。
    Python zipfile 按 cp437 解出来的名字乱码，实测 `encode('cp437').decode('utf-8')`
    能修回正确的中文文件名（本仓库 5 月源材料的 zip 全部如此）。修不回来就原样返回，
    不让编码问题拦掉整个提取。"""
    try:
        return name.encode("cp437").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return name


def _classify_audio(name_hint, date):
    subj = None
    for kw in ("口语", "听力"):
        if kw in name_hint:
            subj = kw
            break
    m = re.search(r"part\s*(\d+)", name_hint, re.I)
    part = f"part{m.group(1)}" if m else ""
    return subj, part


def _dest_audio_name(name_hint, date, ext):
    subj, part = _classify_audio(name_hint, date)
    base = subj or re.sub(r"[^\w\u4e00-\u9fff]+", "", os.path.splitext(name_hint)[0]) or "audio"
    return f"{date} {base}{part}{ext}"


def extract_audio_zip(zip_path, out_dir, date, results):
    with zipfile.ZipFile(zip_path) as zf:
        for info in zf.infolist():
            name = _fix_mojibake_name(info.filename)
            base = os.path.basename(name)
            if not base or base.startswith("._") or "__MACOSX" in name:
                continue
            ext = os.path.splitext(base)[1].lower()
            if ext not in AUDIO_EXT:
                continue
            hint = os.path.basename(zip_path) + name
            out_name = _dest_audio_name(hint, date, ext)
            out_path = os.path.join(out_dir, out_name)
            with zf.open(info) as src, open(out_path, "wb") as dst:
                dst.write(src.read())
            results.append({"src": f"{os.path.basename(zip_path)}::{name}",
                             "out": out_name, "kind": "audio-from-zip"})


def copy_audio_file(path, rel_components, out_dir, date, results):
    ext = os.path.splitext(path)[1].lower()
    hint = "".join(rel_components)
    out_name = _dest_audio_name(hint, date, ext)
    out_path = os.path.join(out_dir, out_name)
    with open(path, "rb") as src, open(out_path, "wb") as dst:
        dst.write(src.read())
    results.append({"src": os.sep.join(rel_components), "out": out_name, "kind": "audio-copy"})


# ── 单套卷驱动 ────────────────────────────────────────────────────────────────
def date_hint_of(dirname):
    """从套目录名里刨出「日期前缀」（用作输出文件名 "<date> 阅读.pdf" 的那一段）。

    自动录入的套名五花八门（"9.12新托福真题" / "5.20" / "9月12日套一"），
    只要能认出 M.D 就用它；认不出就退回目录名本身（照样唯一，只是不好看）。
    """
    m = re.match(r"^\s*(\d{1,2})[.．\-月](\d{1,2})", str(dirname))
    if m:
        return f"{int(m.group(1))}.{int(m.group(2))}"
    return str(dirname).strip() or "set"


def convert_set_dir(src_dir, out_name=None, out_root=OUT_ROOT, date=None):
    """转换**任意**一个套目录（不需要事先登记在 SETS 里）。"""
    if not os.path.isdir(src_dir):
        raise SystemExit(f"源目录不存在: {src_dir}")
    out_name = out_name or os.path.basename(os.path.normpath(src_dir))
    date = date or date_hint_of(out_name)
    return _convert(src_dir, out_name, out_root, date, key=out_name)


def convert_set(key, src_root=None, out_root=OUT_ROOT):
    if key not in SETS:
        raise SystemExit(f"未知卷 key: {key}（可选: {', '.join(SETS)}；"
                         f"任意套目录请改用 --dir <目录>）")
    cfg = SETS[key]
    src_root = src_root or DEFAULT_SRC_ROOT
    src_dir = os.path.join(src_root, cfg["src"])
    if not os.path.isdir(src_dir):
        raise SystemExit(f"源目录不存在: {src_dir}")
    return _convert(src_dir, cfg["out"], out_root, key, key=key)


def _convert(src_dir, out_name, out_root, date, key):
    out_dir = os.path.join(out_root, out_name)
    os.makedirs(out_dir, exist_ok=True)

    files_report = []
    skipped = []

    for root, dirs, names in os.walk(src_dir):
        dirs[:] = [d for d in dirs if d != "__MACOSX"]
        for name in sorted(names):
            if name.startswith("._") or name.startswith("."):
                continue
            path = os.path.join(root, name)
            ext = os.path.splitext(name)[1].lower()
            rel = os.path.relpath(path, src_dir)

            if ext == ".docx":
                subj = subject_of(name)
                if not subj:
                    skipped.append({"src": rel, "reason": "未识别科目关键字"})
                    continue
                out_name = f"{date} {subj}.pdf"
                out_path = os.path.join(out_dir, out_name)
                try:
                    if subj == "答案":
                        n = docx_text_to_pdf(path, out_path)
                        files_report.append({"src": rel, "out": out_name,
                                              "kind": "answer-text-pdf", "paragraphs": n})
                    else:
                        n = docx_images_to_pdf(path, out_path)
                        files_report.append({"src": rel, "out": out_name,
                                              "kind": "stem-image-pdf", "images": n})
                except Exception as e:  # 单份 docx 炸了不该拖垮整套卷的转换
                    skipped.append({"src": rel, "reason": f"转换失败 {type(e).__name__}: {e}"})
            elif ext == ".zip":
                try:
                    extract_audio_zip(path, out_dir, date, files_report)
                except Exception as e:
                    skipped.append({"src": rel, "reason": f"解压失败 {type(e).__name__}: {e}"})
            elif ext in AUDIO_EXT:
                rel_components = os.path.relpath(root, src_dir).split(os.sep) + [name]
                copy_audio_file(path, rel_components, out_dir, date, files_report)
            else:
                skipped.append({"src": rel, "reason": f"未处理的扩展名 {ext}"})

    return {"key": key, "src_dir": src_dir, "out_dir": out_dir,
            "files": files_report, "skipped": skipped}


def report(res):
    print(f"\n{'=' * 74}\n■ {res['key']} → {res['out_dir']}")
    for f in res["files"]:
        extra = f.get("images", f.get("paragraphs", ""))
        print(f"  {f['src'][:40]:42s} → {f['out']:28s} [{f['kind']}] {extra}")
    if res["skipped"]:
        print("  -- 跳过 --")
        for s in res["skipped"]:
            print(f"  ⚠ {s['src']}: {s['reason']}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("key", nargs="?", help=f"卷 key，如 5.10（可选: {', '.join(SETS)}）")
    ap.add_argument("--all", action="store_true", help="转全部 7 套")
    ap.add_argument("--src", default=None, help="覆盖源根目录（默认桌面路径，也可用 REALBANK_SRC 环境变量）")
    ap.add_argument("--dir", default=None,
                    help="直接给套目录（绝对路径，或相对 --src 根的子路径）——不需要登记在 SETS 里")
    ap.add_argument("--out-name", default=None, help="配合 --dir：输出卷名（默认取目录名）")
    ap.add_argument("--out", default=OUT_ROOT, help="覆盖输出根目录")
    ap.add_argument("--json", action="store_true", help="落盘转换报告")
    args = ap.parse_args()

    src_root = args.src or os.environ.get("REALBANK_SRC") or DEFAULT_SRC_ROOT

    all_res = []
    if args.dir:
        src_dir = args.dir if os.path.isabs(args.dir) else os.path.join(src_root, args.dir)
        res = convert_set_dir(src_dir, out_name=args.out_name, out_root=args.out)
        report(res)
        all_res.append(res)
        if args.json:
            os.makedirs(args.out, exist_ok=True)
            p = os.path.join(args.out, f"_convert_{res['key']}.json")
            json.dump(res, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
            print(f"  报告 → {p}")
        return

    keys = list(SETS) if args.all else ([args.key] if args.key else None)
    if not keys:
        ap.error("给个卷 key、用 --all，或用 --dir <套目录>")

    for k in keys:
        res = convert_set(k, src_root=src_root, out_root=args.out)
        report(res)
        all_res.append(res)
        if args.json:
            os.makedirs(os.path.join(args.out), exist_ok=True)
            p = os.path.join(args.out, f"_convert_{k}.json")
            json.dump(res, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
            print(f"  报告 → {p}")

    if len(all_res) > 1:
        print(f"\n共转换 {len(all_res)} 套；输出根目录 → {args.out}")


if __name__ == "__main__":
    main()
